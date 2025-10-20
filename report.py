# report.py
import re
from typing import List, Dict, Tuple
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter


def _extract_group_order(g: str) -> Tuple[int, str]:
    """
    group 문자열에서 선행 숫자를 추출해 정렬 기준을 만든다.
    예) '1학년' -> (1, '1학년'), '4학년' -> (4, '4학년')
        숫자가 없으면 큰 값으로 밀어 뒤로 보냄.
    """
    if g is None:
        return (10**9, "")
    s = str(g)
    m = re.search(r"\d+", s)
    if m:
        return (int(m.group()), s)
    # 숫자가 없으면 문자열 사전순 뒤쪽으로
    return (10**9, s)


def _name_key(name: str) -> str:
    """
    이름 정렬 키. 한국어 가나다 정렬을 완벽히 하지는 않지만,
    일반 문자열 사전순으로 충분히 자연스럽게 정렬됨.
    (필요시 hangul-jamo 분해 로직 추가 가능)
    """
    return str(name or "")


def _auto_group_palette(sorted_groups: List[str]) -> Dict[str, str]:
    """
    group 라벨 리스트(정렬된)를 받아 색상 팔레트를 순서대로 배정.
    색상은 Excel ARGB(hex) 형식으로 지정(FF + RGB).
    팔레트는 파스텔 톤 위주로 구성. 필요시 수정 가능.
    """
    palette = [
        "FFFDE68A",  # 노랑 (1차)
        "FFA7F3D0",  # 민트
        "FFBFDBFE",  # 하늘
        "FFFBCFE8",  # 핑크
        "FFE9D5FF",  # 라일락
        "FFC7D2FE",  # 연보라
        "FFD1FAE5",  # 연민트
        "FFFEE2E2",  # 연빨강
    ]
    color_map: Dict[str, str] = {}
    for i, g in enumerate(sorted_groups):
        color_map[g] = palette[i % len(palette)]
    return color_map


def _best_width_for(texts: List[str], base: int = 12, pad: int = 2) -> int:
    """
    가장 긴 문자열 길이에 맞춰 대략적인 열 너비를 추정.
    (한글은 고정폭이 아니라서 대략치로 충분)
    """
    m = max((len(t or "") for t in texts), default=0)
    return max(base, min(60, m + pad))


def export_team_matrix_excel(
    teams: List[dict],
    filename: str = "output/team_matrix.xlsx",
    name_col: str = "name",
    group_col: str = "group",
    team_label_prefix: str = "Team ",
    show_legend: bool = True,
):
    """
    요구사항:
      - 각 팀을 한 줄(row)로 표현
      - 첫 열(A열)에는 팀 ID(또는 'Team X')
      - 그 뒤로 '이름'만 쭉 가로로 나열
      - 정렬 기준: group 오름차순 → 이름(가나다/사전순)
      - 셀 배경색은 group별 색상. (텍스트로 group/성별 표기 X)

    Parameters
    ----------
    teams : List[dict]
        allocation/balance 결과 구조:
          [{"team_id": 1, "members": [{"id":..., "name":..., "group":..., "gender":...}, ...]}, ...]
    filename : str
        저장할 xlsx 경로
    name_col : str
        이름 컬럼 키(기본 'name', 없으면 id로 대체)
    group_col : str
        그룹 컬럼 키(기본 'group')
    team_label_prefix : str
        첫 열의 팀 라벨 접두사 (예: 'Team ')
    show_legend : bool
        마지막 시트에 group-색상 범례를 출력할지 여부
    """
    # 모든 그룹 수집 + 정렬
    all_groups = set()
    for t in teams:
        for m in t.get("members", []):
            all_groups.add(str(m.get(group_col)))
    groups_sorted = sorted(all_groups, key=lambda g: _extract_group_order(g))

    # group -> 색상 매핑
    group_color = _auto_group_palette(groups_sorted)

    # 워크북 생성
    wb = Workbook()
    ws = wb.active
    ws.title = "Teams"

    # 헤더(선택적): 팀 라벨만 굵게
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # 각 팀 한 줄씩 작성
    max_name_count = 0
    for row_idx, team in enumerate(teams, start=1):
        team_id = team.get("team_id")
        members = team.get("members", [])

        # 정렬: (group asc, name asc)
        members_sorted = sorted(
            members,
            key=lambda m: (_extract_group_order(str(m.get(group_col))), _name_key(m.get(name_col) or m.get("id")))
        )

        # A열: 팀 라벨
        ws.cell(row=row_idx, column=1, value=f"{team_label_prefix}{team_id}").font = header_font

        # B열 이후: 이름만 쓰고, 배경색은 group 색상
        for j, m in enumerate(members_sorted, start=2):
            display_name = str(m.get(name_col) or m.get("id") or "")
            g = str(m.get(group_col))
            cell = ws.cell(row=row_idx, column=j, value=display_name)
            cell.alignment = center
            # 배경색 채우기
            fill_color = group_color.get(g)
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)

        max_name_count = max(max_name_count, len(members_sorted))

    # 열 너비 조정 (이름 열 기준)
    # 1열은 'Team X' 라벨 위한 너비, 2열~N열은 이름 길이 기반
    ws.column_dimensions["A"].width = 12
    # 수집해서 대략의 너비 계산
    all_names = []
    for t in teams:
        for m in t.get("members", []):
            all_names.append(str(m.get(name_col) or m.get("id") or ""))
    width = _best_width_for(all_names, base=12, pad=2)

    for col in range(2, max_name_count + 2):
        ws.column_dimensions[get_column_letter(col)].width = width

    # 보기 좋게 첫 행 고정(원하면) — 여기서는 헤더행이 따로 없으므로 생략
    # ws.freeze_panes = "B2"

    # 범례 시트 (옵션)
    if show_legend:
        ws2 = wb.create_sheet("Legend")
        ws2.append(["Group", "Color"])
        ws2["A1"].font = header_font
        ws2["B1"].font = header_font
        for i, g in enumerate(groups_sorted, start=2):
            ws2.cell(row=i, column=1, value=g)
            c = ws2.cell(row=i, column=2, value="")
            fill_color = group_color.get(g)
            if fill_color:
                c.fill = PatternFill("solid", fgColor=fill_color)
        ws2.column_dimensions["A"].width = 16
        ws2.column_dimensions["B"].width = 16

    # 저장
    # 경로상의 폴더가 없으면 에러가 날 수 있으니, 사용자가 미리 만들어두는 것을 권장
    wb.save(filename)
    print(f"✅ Report saved to: {filename}")


# CSV 버전도 원하면 아래 간단히 제공(색칠 불가, 이름만 가로로)
def export_team_matrix_csv(
    teams: List[dict],
    filename: str = "output/team_matrix.csv",
    name_col: str = "name",
    group_col: str = "group",
    team_label_prefix: str = "Team "
):
    """
    팀별 한 줄, 이름 가로 나열(정렬 동일). 색칠은 CSV 특성상 불가.
    """
    rows = []
    for team in teams:
        team_id = team.get("team_id")
        members = team.get("members", [])
        members_sorted = sorted(
            members,
            key=lambda m: (_extract_group_order(str(m.get(group_col))), _name_key(m.get(name_col) or m.get("id")))
        )
        # 한 행: [Team X, name1, name2, ...]
        row = [f"{team_label_prefix}{team_id}"] + [str(m.get(name_col) or m.get("id") or "") for m in members_sorted]
        rows.append(row)

    # 가변 열 폭을 가진 CSV를 DataFrame으로 저장
    max_len = max((len(r) for r in rows), default=1)
    # pad rows to same length
    norm = [r + [""] * (max_len - len(r)) for r in rows]
    df = pd.DataFrame(norm)
    df.to_csv(filename, index=False, header=False, encoding="utf-8-sig")
    print(f"✅ CSV saved to: {filename}")